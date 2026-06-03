from __future__ import annotations

from pathlib import Path

import pytest

from icm_engine.excel import read_xlsx_rows, write_xlsx

FIXTURES = Path(__file__).parent / "fixtures" / "excel"


def test_read_clean_transactions() -> None:
    headers, rows = read_xlsx_rows(FIXTURES / "clean_transactions.xlsx")

    assert headers == ["id", "payee_id", "deal_id", "period", "amount", "product", "close_date"]
    assert len(rows) == 3
    assert rows[0]["id"] == "T001"
    assert rows[0]["payee_id"] == "P001"
    assert rows[0]["amount"] == "1500.00"
    assert rows[0]["close_date"] == "2026-04-15"
    assert rows[2]["product"] == ""  # empty cell


def test_read_messy_transactions() -> None:
    headers, rows = read_xlsx_rows(FIXTURES / "messy_transactions.xlsx")

    # Header should be detected at row 3 (row 1 is merged title)
    assert "transaction_id" in headers
    assert "rep_name" in headers
    assert "tcv" in headers

    # Should have exactly 3 data rows (totals row skipped)
    assert len(rows) == 3

    # Dollar signs preserved as-is by excel reader (strip_currency in mapping handles cleanup)
    first_amount = rows[0]["tcv"]
    assert "$" in first_amount

    # MM/DD date preserved as string
    assert "/" in rows[0]["close_date"]


def test_header_detection_messy() -> None:
    headers, rows = read_xlsx_rows(FIXTURES / "messy_transactions.xlsx")

    # Verify header auto-detection found row 3 (actual headers), not row 1 (merged title)
    assert all(h == h.lower() for h in headers)
    assert all(" " not in h for h in headers)

    # Row 1 merged title repeated across all cells: "Q1 2026 Sales Report"
    # Row 3 real headers are all unique. Unique-count tiebreaker picks row 3.
    assert len(rows) == 3
    for row in rows:
        assert any(v for v in row.values()), f"Row should not be all blank: {row}"


def test_totals_row_skipped() -> None:
    headers, rows = read_xlsx_rows(FIXTURES / "messy_transactions.xlsx")

    # The "Total" row (row 7 in sheet) should not appear in data rows
    for row in rows:
        first_val = list(row.values())[0] if row.values() else ""
        assert first_val.lower() not in ("total", "grand total", "totals")


def test_merged_cell_propagation() -> None:
    _, rows = read_xlsx_rows(FIXTURES / "messy_transactions.xlsx")

    # After unmerging, the merged "Total" row has "Total" in all cells A7:F7
    # That triggers _drop_total_rows, leaving 3 data rows
    assert len(rows) == 3


def test_multi_sheet_reads_first_with_data() -> None:
    headers, rows = read_xlsx_rows(FIXTURES / "multi_sheet.xlsx")

    # Picks first non-empty sheet ("Transaction Data"), reads its headers + row
    assert "id" in headers
    assert "payee_id" in headers
    assert len(rows) == 1
    assert rows[0]["id"] == "T001"


def test_write_then_read_decimal_preservation(tmp_path: Path) -> None:
    """Round-trip: write commissions then read them back, verify Decimal precision."""
    rows = [
        {"transaction_id": "T001", "payee_id": "P001", "period": "2026-04",
         "rule_id": "flat", "base_amount": "1500.00", "rate": "0.05",
         "commission_amount": "75.00", "notes": ""},
        {"transaction_id": "T002", "payee_id": "P001", "period": "2026-04",
         "rule_id": "flat", "base_amount": "2500.00", "rate": "0.05",
         "commission_amount": "125.00", "notes": ""},
    ]
    out_path = tmp_path / "test_output.xlsx"
    write_xlsx(out_path, {"commissions": rows})

    assert out_path.exists()
    headers, read_rows = read_xlsx_rows(out_path)

    assert "commission_amount" in headers
    assert len(read_rows) == 2
    assert read_rows[0]["commission_amount"] == "75.00"
    assert read_rows[1]["commission_amount"] == "125.00"


def test_write_then_read_summary(tmp_path: Path) -> None:
    rows = [
        {"payee_id": "P001", "period": "2026-04", "total_commission": "200.00"},
        {"payee_id": "P002", "period": "2026-04", "total_commission": "50.00"},
    ]
    out_path = tmp_path / "summary.xlsx"
    write_xlsx(out_path, {"summary": rows})

    headers, read_rows = read_xlsx_rows(out_path)
    assert len(read_rows) == 2
    assert read_rows[0]["payee_id"] == "P001"
    assert read_rows[0]["total_commission"] == "200.00"


def test_write_multiple_sheets(tmp_path: Path) -> None:
    out_path = tmp_path / "multi_out.xlsx"
    write_xlsx(out_path, {
        "commissions": [
            {"transaction_id": "T001", "commission_amount": "75.00"},
        ],
        "summary": [
            {"payee_id": "P001", "total_commission": "75.00"},
        ],
    })
    assert out_path.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    assert "commissions" in wb.sheetnames
    assert "summary" in wb.sheetnames


def test_empty_rows_input() -> None:
    """Non-existent file raises FileNotFoundError."""
    with pytest.raises(FileNotFoundError):
        read_xlsx_rows(Path("nonexistent.xlsx"))


def test_money_columns_formatted(tmp_path: Path) -> None:
    """Verify money columns get dollar number format."""
    out_path = tmp_path / "money.xlsx"
    write_xlsx(out_path, {
        "commissions": [
            {"commission_amount": "1234.5678", "rate": "0.0500", "notes": "test"},
        ],
    })

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    ws = wb["commissions"]

    # commission_amount column (col 1) should have number format
    comm_cell = ws.cell(row=2, column=1)
    assert comm_cell.number_format == "$#,##0.0000"

    # Header should be bold
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold


def test_money_columns_detected_by_suffix(tmp_path: Path) -> None:
    """Columns ending in _amount should get money formatting."""
    out_path = tmp_path / "suffix.xlsx"
    write_xlsx(out_path, {
        "sheet1": [{"base_amount": "100.50"}],
    })

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    ws = wb["sheet1"]

    cell = ws.cell(row=2, column=1)
    assert cell.number_format == "$#,##0.0000"


def test_frozen_panes(tmp_path: Path) -> None:
    out_path = tmp_path / "frozen.xlsx"
    write_xlsx(out_path, {"data": [{"col": "val"}]})

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    ws = wb["data"]
    assert ws.freeze_panes == "A2"
